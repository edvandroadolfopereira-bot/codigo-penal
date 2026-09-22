# -*- coding: utf-8 -*-
"""Poe a BASE ANALITICA dentro da planilha entregue, e corrige a numeracao das abas.

POR QUE EXISTE, e o achado e' de 18/09/2026: a auditoria metodologica de 16/09/2026
concluiu que os 130 eventos e os 432 dispositivos "nao sao reproduziveis", porque a
planilha nao os continha. As duas bases EXISTIAM no pacote (cp-eventos-de-severidade.csv
e cp-severidade-por-artigo.csv) e nunca entraram na planilha. Pior: a aba chamada
"3 Eventos de severidade" trazia as 1.116 anotacoes de alteracao, e nao os eventos. O
nome errado da aba produziu a conclusao de base ausente.

O QUE FAZ:
  1. renomeia a aba 3 para o que ela e' (anotacoes de alteracao por norma);
  2. acrescenta a aba dos 130 eventos, uma linha por evento, com materia (Titulo do
     Codigo), origem da iniciativa, fonte do texto e a REGRA aplicada por extenso;
  3. acrescenta a aba dos 432 dispositivos, 1940 contra o texto vigente;
  4. renumera abas e tabelas em sequencia, o que desfaz as duas abas "22";
  5. atualiza o indice.

A REGUA e' a do proprio trabalho, em severidade_por_evento_cp.py: o teto decide; havendo
empate, decide o piso. Ela e' RECALCULADA aqui sobre os numeros de cada linha, e a
gravacao e' RECUSADA se um unico evento divergir da direcao gravada na base.

Materia e origem saem das MESMAS funcoes que produziram os testes do trabalho
(apurar_materia_das_normas.py e apurar_autoria_e_bancada.py), para a planilha nao ter
uma segunda verdade.

Grava de forma atomica: monta em memoria, confere, e so' entao substitui.
"""

import collections
import csv
import io
import json
import os
import re
import sqlite3
import sys

import openpyxl
from copy import copy

sys.stdout.reconfigure(encoding="utf-8")

R = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, R)
import apurar_materia_das_normas as MAT  # noqa: E402
import apurar_autoria_e_bancada as AUT  # noqa: E402

PLANILHA = os.path.join(R, "ENTREGA-ACADEMICA", "03-BANCO-DE-DADOS-CODIGO-PENAL.xlsx")
EVENTOS = os.path.join(R, "cp-eventos-de-severidade.csv")
DISPOSITIVOS = os.path.join(R, "cp-severidade-por-artigo.csv")
TEXTO = os.path.join(R, "cp-compilado-texto.txt")
BASE = os.path.join(R, "BASE-CODIGO-PENAL.sqlite")
TEXTOS_NORMAS = os.path.join(R, "textos-das-normas")

ABA_ANOTACOES_ANTIGA = "3 Eventos de severidade"
TRAVESSAO = chr(0x2014)
# Sinais convencionais das Normas de Apresentacao Tabular do IBGE, 3. ed., 1993,
# item 4.8.1, os mesmos declarados no indice da planilha. Celula vazia NAO e' sinal:
# o leitor nao sabe se e' zero, se falta o dado ou se nao se aplica.
NAO_SE_APLICA = ".."        # ex.: pena anterior de artigo que nao existia em 1940
NAO_DISPONIVEL = "..."      # dado que deveria existir e nao foi obtido


def ler_csv(p):
    with io.open(p, encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f, delimiter=";"))


def num(v):
    v = (v or "").strip()
    return int(v) if v else None


def regra(e):
    """Recalcula a direcao pela regua do trabalho e devolve (direcao, frase)."""
    pa, ta = num(e["piso_antes"]), num(e["teto_antes"])
    pd, td = num(e["piso_depois"]), num(e["teto_depois"])
    if ta is None and td is not None:
        if e["estado_anterior"] == "artigo inexistente em 1940":
            return "cria pena", "tipo novo: artigo criado já com pena cominada"
        return "cria pena", "artigo sem pena cominada no estado anterior passou a ter pena"
    if ta is not None and td is None:
        return "retira pena", "pena cominada retirada"
    if td > ta:
        return "agrava", "teto subiu de %d para %d meses" % (ta, td)
    if td < ta:
        return "abranda", "teto caiu de %d para %d meses" % (ta, td)
    if pd > pa:
        return "agrava", "teto igual (%d meses); piso subiu de %d para %d meses" % (ta, pa, pd)
    if pd < pa:
        return "abranda", "teto igual (%d meses); piso caiu de %d para %d meses" % (ta, pa, pd)
    return "mantem", "teto e piso iguais (%d a %d meses)" % (pa, ta)


def fonte_dos_textos():
    """Qual corpus deu o texto de cada norma, pela mesma preferencia do gerador:
    o Planalto vence quando existe; senao, o texto do Congresso."""
    fonte = {}
    for nome, rotulo in (("texto-integral-normas-cp.jsonl", "texto da norma, Congresso Nacional"),
                         ("texto-integral-planalto.jsonl", "texto da norma, Presidência da República")):
        p = os.path.join(TEXTOS_NORMAS, nome)
        if not os.path.exists(p):
            continue
        for linha in io.open(p, encoding="utf-8"):
            linha = linha.strip()
            if linha:
                d = json.loads(linha)
                fonte[(str(d["numero"]).lstrip("0"), str(d["ano"]))] = rotulo
    return fonte


def fonte_anterior(estado):
    if estado == "texto original de 1940":
        return "publicação original de 1940, Legin da Câmara dos Deputados"
    if estado == "artigo inexistente em 1940":
        return NAO_SE_APLICA
    return "texto da Lei %s" % estado


def materia_por_artigo():
    texto = io.open(TEXTO, encoding="utf-8").read()
    est = MAT.faixas_de_artigo(texto, MAT.estrutura(texto))
    mapa, _ = MAT.mapa_artigo_para_titulo(est)
    return mapa


def origem_por_norma():
    c = sqlite3.connect(BASE)
    c.row_factory = sqlite3.Row
    por = collections.defaultdict(list)
    for r in c.execute("SELECT norma, autor, tipo_de_autor FROM camara_autores"):
        k = AUT.chave(r["norma"])
        if k:
            por[k].append((r["tipo_de_autor"], r["autor"]))
    c.close()
    return {k: AUT.origem_da_iniciativa(v) for k, v in por.items()}


def montar_eventos():
    eventos = ler_csv(EVENTOS)
    mapa = materia_por_artigo()
    origem = origem_por_norma()
    fonte = fonte_dos_textos()
    linhas, divergencias = [], []
    for i, e in enumerate(eventos, 1):
        d, frase = regra(e)
        if d != e["direcao"]:
            divergencias.append((i, e["dispositivo"], e["norma"], e["direcao"], d))
        t = mapa.get(MAT.artigo_do_dispositivo(e["dispositivo"]))
        num_n, ano_n = e["norma"].split("/")
        linhas.append([
            "E%03d" % i, int(e["ano"]), "Lei " + e["norma"], e["dispositivo"], e["parte"],
            ("Título %s, %s" % (t["titulo"], t["nome"])) if t else "...",
            e["estado_anterior"],
            num(e["piso_antes"]), num(e["teto_antes"]),
            num(e["piso_depois"]), num(e["teto_depois"]),
            num(e["variacao_do_teto"]),
            e["direcao"],
            "sim" if e["direcao"] in ("agrava", "cria pena") else "não",
            frase,
            origem.get(AUT.chave(e["norma"]), "sem dado de autoria"),
            fonte_anterior(e["estado_anterior"]),
            fonte.get((num_n.lstrip("0"), ano_n), "..."),
        ])
    cab = ["id_evento", "ano", "norma", "dispositivo", "parte", "materia_titulo_do_codigo",
           "estado_anterior", "piso_antes_meses", "teto_antes_meses", "piso_depois_meses",
           "teto_depois_meses", "variacao_do_teto_meses", "direcao", "endurece",
           "regra_aplicada", "origem_da_iniciativa", "fonte_do_estado_anterior",
           "fonte_do_estado_posterior"]
    return cab, linhas, divergencias


def montar_dispositivos():
    ds = ler_csv(DISPOSITIVOS)
    cab = ["id_dispositivo"] + list(ds[0].keys())
    linhas = []
    for i, d in enumerate(ds, 1):
        linha = ["D%03d" % i]
        for k, v in d.items():
            linha.append(num(v) if k.endswith("_meses") and (v or "").strip().lstrip("-").isdigit() else v)
        linhas.append(linha)
    return cab, linhas


def escrever_aba(ws, modelo, titulo, cab, linhas, fonte, notas):
    ws.cell(1, 1, titulo).font = copy(modelo.cell(1, 1).font)
    for j, h in enumerate(cab, 1):
        c = ws.cell(3, j, h)
        c.font = copy(modelo.cell(3, 1).font)
    for i, linha in enumerate(linhas, 4):
        for j, v in enumerate(linha, 1):
            # Nas duas bases, campo vazio so' ocorre onde a pena nao existe de um dos
            # lados (tipo novo, pena retirada, artigo sem faixa em meses), ou seja,
            # onde o numero nao se aplica. Por isso vazio vira "..", e nunca branco.
            ws.cell(i, j, NAO_SE_APLICA if v is None or v == "" else v)
    r = 4 + len(linhas) + 1
    ws.cell(r, 1, "Fonte: " + fonte)
    for n in notas:
        r += 1
        ws.cell(r, 1, "Nota: " + n)
    ws.freeze_panes = "A4"
    for j, h in enumerate(cab, 1):
        larg = max([len(str(h))] + [len(str(l[j - 1])) for l in linhas if l[j - 1] is not None])
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = min(max(10, larg + 2), 60)


def main():
    for p in (PLANILHA, EVENTOS, DISPOSITIVOS, TEXTO, BASE):
        if not os.path.exists(p):
            print("FALHA REAL: falta " + p)
            return 1

    cab_e, eventos, diverg = montar_eventos()
    if diverg:
        print("RECUSADO: %d evento(s) divergem da regua do trabalho, nada foi gravado:" % len(diverg))
        for d in diverg:
            print("   ", d)
        return 1
    cab_d, disps = montar_dispositivos()

    cont = collections.Counter(l[12] for l in eventos)
    endurece = sum(1 for l in eventos if l[13] == "sim")
    sem_materia = [l[0] for l in eventos if l[5] == "..."]
    sem_origem = [l[0] for l in eventos if l[15] == "sem dado de autoria"]
    sem_fonte = [l[0] for l in eventos if l[17] == "..."]
    normas = len(set(l[2] for l in eventos))
    dcont = collections.Counter(l[cab_d.index("direcao")] for l in disps)

    wb = openpyxl.load_workbook(PLANILHA)
    if ABA_ANOTACOES_ANTIGA not in wb.sheetnames:
        print("RECUSADO: a aba '%s' nao existe; a planilha ja' foi corrigida ou e' outra versao."
              % ABA_ANOTACOES_ANTIGA)
        return 1
    modelo = wb[ABA_ANOTACOES_ANTIGA]
    pos = wb.sheetnames.index(ABA_ANOTACOES_ANTIGA)
    nomes_antigos = list(wb.sheetnames)

    ws_e = wb.create_sheet("tmp_eventos", pos + 1)
    ws_d = wb.create_sheet("tmp_dispositivos", pos + 2)
    escrever_aba(
        ws_e, modelo,
        "Tabela X - Eventos datados de severidade, um por transição de pena - Brasil - 1968-2026",
        cab_e, eventos,
        "Cadeia de penas por dispositivo, a partir da publicação original de 1940 (Legin da Câmara) "
        "e do texto de cada norma alteradora (Presidência da República e Congresso Nacional)",
        ["régua: o teto da pena cominada decide; havendo empate, decide o piso. Valores em meses",
         "evento é cada transição datada de pena de um artigo; a mesma lei pode produzir vários eventos",
         "a série é rala antes de 1990 porque o texto compilado exibe apenas a anotação sobrevivente",
         "alteração que não escreve pena nova não produz evento; causa de aumento e de diminuição não entram",
         "matéria é o Título do próprio Código Penal que contém o artigo; não é categoria do autor"])
    escrever_aba(
        ws_d, modelo,
        "Tabela X - Dispositivos do Código Penal, pena de 1940 contra pena vigente - Brasil - 1940-2026",
        cab_d, disps,
        "Publicação original de 1940 (Legin da Câmara) e texto compilado da Presidência da República",
        ["régua: o teto da pena cominada decide; havendo empate, decide o piso. Valores em meses",
         "efeito líquido entre 1940 e o texto vigente; não traz data, que está na aba dos eventos"])

    # renumeracao sequencial das abas
    base_nome = {}
    for ws in wb.worksheets:
        base_nome[ws.title] = re.sub(r"^\d+\s+", "", ws.title)
    base_nome[ABA_ANOTACOES_ANTIGA] = "Anotacoes de alteracao"
    base_nome["tmp_eventos"] = "Eventos de severidade"
    base_nome["tmp_dispositivos"] = "Dispositivos 1940 e hoje"
    novo = {}
    for i, ws in enumerate(wb.worksheets, 1):
        novo[ws.title] = ("%d %s" % (i, base_nome[ws.title]))[:31]
    for ws in wb.worksheets:
        ws.title = novo[ws.title]

    # renumeracao sequencial das tabelas, na linha 1 de cada aba de dados
    n_tab = 0
    for ws in wb.worksheets[1:]:
        v = ws.cell(1, 1).value
        if isinstance(v, str) and re.match(r"^Tabela (\d+|X) - ", v):
            n_tab += 1
            ws.cell(1, 1, re.sub(r"^Tabela (\d+|X)", "Tabela %d" % n_tab, v))

    # indice
    ind = wb.worksheets[0]
    linha_ab = None
    for r in range(1, ind.max_row + 1):
        if ind.cell(r, 1).value == "Aba":
            linha_ab = r
            break
    if linha_ab is None:
        print("RECUSADO: cabecalho 'Aba' do indice nao encontrado")
        return 1
    primeira = linha_ab + 1
    antigas = []
    r = primeira
    while ind.cell(r, 1).value in nomes_antigos:
        antigas.append([ind.cell(r, c).value for c in range(1, 6)])
        r += 1
    if len(antigas) != len(nomes_antigos) - 1:
        print("RECUSADO: o indice lista %d abas e a planilha tem %d" % (len(antigas), len(nomes_antigos) - 1))
        return 1
    ind.insert_rows(primeira + len(antigas), 2)
    novas = []
    for a in antigas:
        a = list(a)
        a[0] = novo[a[0]]
        novas.append(a)
        if a[0] == novo[ABA_ANOTACOES_ANTIGA]:
            novas.append([novo["tmp_eventos"], None, len(eventos), "-",
                          "Cadeia de penas por dispositivo, 1940 e normas alteradoras"])
            novas.append([novo["tmp_dispositivos"], None, len(disps), "-",
                          "Publicação original de 1940 e texto compilado da Presidência"])
    for i, a in enumerate(novas):
        for c, v in enumerate(a, 1):
            ind.cell(primeira + i, c, v)
            ind.cell(primeira + i, c).font = copy(ind.cell(primeira, c).font)
        titulo_aba = wb[a[0]].cell(1, 1).value
        if isinstance(titulo_aba, str) and titulo_aba.startswith("Tabela"):
            ind.cell(primeira + i, 2, titulo_aba)

    # conferencia antes de gravar. O travessao se confere no que ESTE script escreve
    # (indice e abas novas). Nas abas de fonte ele vem transcrito da origem, como as
    # ementas do STF, e trocar caractere de ementa seria adulterar citacao: la' ele
    # e' so' contado e declarado.
    escritas = {wb.worksheets[0].title, novo["tmp_eventos"], novo["tmp_dispositivos"]}
    em_fonte = collections.Counter()
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if isinstance(v, str) and TRAVESSAO in v:
                    if ws.title in escritas:
                        print("RECUSADO: travessao na aba " + ws.title)
                        return 1
                    em_fonte[ws.title] += 1
    if len(set(wb.sheetnames)) != len(wb.sheetnames):
        print("RECUSADO: nomes de aba repetidos")
        return 1

    tmp = PLANILHA + ".tmp.xlsx"
    wb.save(tmp)
    wb2 = openpyxl.load_workbook(tmp)
    ok_e = sum(1 for _ in wb2[novo["tmp_eventos"]].iter_rows(min_row=4, values_only=True)
               if _[0] and str(_[0]).startswith("E0") or (_[0] and str(_[0]).startswith("E1")))
    ok_d = sum(1 for _ in wb2[novo["tmp_dispositivos"]].iter_rows(min_row=4, values_only=True)
               if _[0] and re.match(r"^D\d{3}$", str(_[0])))
    if ok_e != 130 or ok_d != 432:
        os.remove(tmp)
        print("RECUSADO: releitura trouxe %d eventos e %d dispositivos" % (ok_e, ok_d))
        return 1
    os.replace(tmp, PLANILHA)

    print("GRAVADO: " + PLANILHA)
    print("abas: %d  (antes %d)" % (len(wb2.sheetnames), len(nomes_antigos)))
    for a, b in novo.items():
        if a != b:
            print("   %-32s -> %s" % (a, b))
    print("eventos relidos: %d | dispositivos relidos: %d | normas com evento: %d" % (ok_e, ok_d, normas))
    print("eventos por direcao: %s | endurecem (agrava + cria pena): %d" % (dict(cont), endurece))
    print("dispositivos por direcao: %s" % dict(dcont))
    print("regua recalculada: 130 de 130 conferem com a direcao gravada")
    print("travessao transcrito da fonte (celulas, nao alterado): %s" % (dict(em_fonte) or "nenhum"))
    print("sem materia: %s | sem origem: %s | sem fonte do texto: %s"
          % (sem_materia or "nenhum", sem_origem or "nenhum", sem_fonte or "nenhum"))
    return 0


if __name__ == "__main__":
    sys.exit(main())
