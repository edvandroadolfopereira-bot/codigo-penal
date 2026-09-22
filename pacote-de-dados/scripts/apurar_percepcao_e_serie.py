# -*- coding: utf-8 -*-
# A PERCEPCAO DE INSEGURANCA, e a SERIE DE MORTES VIOLENTAS corrigida.
#
# POR QUE ELE EXISTE, e a razao e' um erro meu: em 30/08/2026 declarei os dois itens
# impossiveis apoiado em sondagem de rede, e os dois eram resolviveis com material ja'
# no disco. O acervo do Forum Brasileiro de Seguranca Publica tem 578 arquivos e 1,9
# GB, e eu nunca o abri.
#
# O ERRO DE INSTRUMENTO QUE O ESCONDEU: a base grava a raiz do acervo como `forum`, e
# eu testei a existencia dos arquivos sem resolver o caminho absoluto. Deu ausente em
# 78 arquivos, e eu li aquilo como ausencia real. E' a regra de 27/07/2026: nunca
# concluir ausencia sem localizar por conteudo.
#
# O QUE ELE MEDE:
#   A. a serie de Mortes Violentas Intencionais, por EDICAO do Anuario, o que expoe
#      tanto o alcance real quanto a retificacao entre edicoes;
#   B. a percepcao de inseguranca, comparada entre 2017 e 2022, do estudo do proprio
#      Forum, que e' a hipotese do titulo deste trabalho nunca antes medida.
#
# O QUE ELE NAO FAZ: nao constroi serie anual de percepcao. Sao DOIS pontos, 2017 e
# 2022, e dois pontos comparam, mas nao formam serie. Isso se declara.

import collections
import io
import json
import os
import re
import sys

sys.stdout.reconfigure(encoding="utf-8")

R = os.path.dirname(os.path.abspath(__file__))
# A Biblioteca mora em arvores diferentes conforme a maquina, e caminho fixo envelhece:
# medido em 22/09/2026, neste PC ela esta' em OneDrive\.claude\ e nao em OneDrive\.
# Vale a primeira que existe; nenhuma existindo, a bancada acusa.
_CANDIDATAS = [os.path.join(os.environ.get("USERPROFILE", ""), "OneDrive", ".claude",
                            "Biblioteca-Academica-Direito", "forum-seguranca-publica"),
               os.path.join(os.environ.get("USERPROFILE", ""), "OneDrive",
                            "Biblioteca-Academica-Direito", "forum-seguranca-publica")]
ACERVO = next((c for c in _CANDIDATAS if os.path.isdir(c)), _CANDIDATAS[0])
SAIDA = os.path.join(R, "analise-cientifica", "percepcao-e-serie.json")

# Cada par declara a PERGUNTA e as duas medidas, e o trecho de onde cada uma sai. O
# valor nao e' digitado: e' extraido do texto do estudo, e o trecho fica gravado.
PARES = [
    {"medo_de": "morrer assassinado",
     "padrao": r"(\d{2},\d)% t[êe]m este medo em 2022; em 2017 eram (\d{2},\d)%"},
    {"medo_de": "ter os dados pessoais divulgados na internet",
     "padrao": r"em 2022, (\d{2},\d)% t[êe]m medo de\s*ter seus dados pessoais divulgados na internet, enquanto em\s*2017, (\d{2},\d)%"},
    {"medo_de": "sofrer violencia por parte das policias militares",
     "padrao": r"est[áa] em (\d{2},\d)% dos entrevistados, contra (\d{2},\d)% em 2017"},
]

# medidas de 2022 sem par em 2017, que entram como ponto unico e assim se declaram
PONTO_UNICO = [
    {"o_que": "medo da violencia politica, ser agredido por escolha politica ou partidaria",
     "padrao": r"medo da viol[êe]ncia pol[íi]tica aparece em (\d{2},\d)% dos entrevistados"},
    {"o_que": "medo de ser vitima de grupos armados",
     # o texto do PDF quebra a linha no meio da expressao, e por isso o padrao tem que
     # atravessar espaco em branco qualquer, inclusive quebra
     "padrao": r"grupos armados\s*\(tra-?\s*ficantes,\s*mil[íi]cias e pistoleiros\)\s*est[áa] em\s*(\d{2},\d)%"},
    {"o_que": "muito medo de ser vitima de grupos armados",
     "padrao": r"sendo que (\d{2},\d)% dos entrevistados afirmam ter muito medo"},
]

FRASE_CHAVE = r"queda nas mortes violentas\s*intencionais.{0,120}?n[ãa]o foi percebida pelos entrevistados"


def texto_dos_estudos():
    """Le o texto dos PDF de percepcao do acervo, do primeiro ao ultimo, sem amostra."""
    import pymupdf
    alvos = []
    for d, _, arq in os.walk(ACERVO):
        for a in arq:
            n = a.lower()
            if n.endswith(".pdf") and any(t in n for t in
                                          ("medo", "eleic", "eleiç", "autoritarismo",
                                           "democracia", "vitimiza")):
                alvos.append(os.path.join(d, a))
    partes, lidos = [], []
    for p in sorted(alvos):
        try:
            doc = pymupdf.open(p)
            partes.append("\n".join(doc[i].get_text() for i in range(doc.page_count)))
            lidos.append({"arquivo": os.path.basename(p), "paginas": doc.page_count,
                          "bytes": os.path.getsize(p)})
            doc.close()
        except Exception as e:
            lidos.append({"arquivo": os.path.basename(p), "falha": type(e).__name__})
    return "\n".join(partes), lidos


def serie_mvi():
    """A serie do Brasil, por edicao do Anuario, para expor a retificacao entre elas."""
    import openpyxl
    plan, vistos = [], set()
    for d, _, arq in os.walk(ACERVO):
        for a in arq:
            if a.lower().endswith(".xlsx") and "anuario" in a.lower() and a not in vistos:
                vistos.add(a)
                plan.append(os.path.join(d, a))
    serie, notas = {}, set()
    for p in sorted(plan):
        nome = os.path.basename(p)
        try:
            wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        except Exception:
            continue
        for aba in wb.sheetnames:
            cab, dentro = None, False
            for linha in wb[aba].iter_rows(values_only=True):
                vals = [("" if x is None else str(x).strip()) for x in linha]
                txt = " ".join(vals)
                if re.search(r"S[ée]rie hist[oó]rica das Mortes Violentas", txt, re.I):
                    dentro = True
                if not dentro:
                    continue
                if re.search(r"s[óo] passou a ser calculada", txt, re.I):
                    notas.add(re.sub(r"\s+", " ", txt.split("|")[0]).strip())
                anos = [(i, int(v.replace(".0", ""))) for i, v in enumerate(vals)
                        if re.fullmatch(r"(19[7-9]\d|20[0-2]\d)(\.0)?", v)]
                if len(anos) >= 5 and cab is None:
                    cab = anos
                    continue
                if cab and vals and vals[0].strip().lower() == "brasil":
                    for i, ano in cab:
                        if i < len(vals) and vals[i]:
                            try:
                                v = float(vals[i])
                            except ValueError:
                                continue
                            if v > 1000:
                                serie.setdefault(ano, {})[nome] = int(v)
                    dentro, cab = False, None
        wb.close()
    return serie, sorted(notas), sorted(os.path.basename(x) for x in plan)


def apurar():
    if not os.path.isdir(ACERVO):
        print("FALHA REAL: acervo do Forum ausente em " + ACERVO)
        return 1

    serie, notas, edicoes = serie_mvi()
    t, lidos = texto_dos_estudos()

    pares, nao_lidos = [], []
    for p in PARES:
        m = re.search(p["padrao"], t, re.I | re.S)
        if not m:
            nao_lidos.append(p["medo_de"])
            continue
        pares.append({"medo_de": p["medo_de"],
                      "em_2022": float(m.group(1).replace(",", ".")),
                      "em_2017": float(m.group(2).replace(",", ".")),
                      "trecho": re.sub(r"\s+", " ", m.group(0))})
    unicos = []
    for p in PONTO_UNICO:
        m = re.search(p["padrao"], t, re.I | re.S)
        if not m:
            nao_lidos.append(p["o_que"])
            continue
        unicos.append({"o_que": p["o_que"], "em_2022": float(m.group(1).replace(",", ".")),
                       "trecho": re.sub(r"\s+", " ", m.group(0))})
    achou_frase = re.search(FRASE_CHAVE, re.sub(r"\s+", " ", t), re.I)

    anos = sorted(serie)
    retificados = [a for a in anos if len(set(serie[a].values())) > 1]

    res = {
        "gerado_por": "apurar_percepcao_e_serie.py",
        "acervo": "Forum Brasileiro de Seguranca Publica, copia local",
        "A_serie_de_mortes_violentas_intencionais": {
            "alcance": {"de": min(anos), "a": max(anos), "anos": max(anos) - min(anos) + 1},
            "edicoes_lidas": edicoes,
            "por_ano": {str(a): {"valores_publicados": sorted(set(serie[a].values())),
                                 "edicoes_que_publicam": len(serie[a])} for a in anos},
            "anos_com_valor_retificado_entre_edicoes": retificados,
            "notas_da_propria_fonte": notas,
        },
        "B_percepcao_de_inseguranca": {
            "estudos_lidos": lidos,
            "comparacao_2017_contra_2022": pares,
            "medidas_de_2022_sem_par_em_2017": unicos,
            "a_fonte_afirma_que_a_queda_da_violencia_nao_foi_percebida": bool(achou_frase),
            "trecho_da_afirmacao": re.sub(r"\s+", " ", achou_frase.group(0)) if achou_frase else None,
            "o_que_isto_NAO_e": ("serie anual de percepcao. Sao dois pontos, 2017 e 2022, e "
                                 "dois pontos comparam, mas nao formam serie"),
        },
        "medidas_declaradas_e_nao_lidas": nao_lidos,
    }
    os.makedirs(os.path.dirname(SAIDA), exist_ok=True)
    io.open(SAIDA, "w", encoding="utf-8", newline="").write(
        json.dumps(res, ensure_ascii=False, indent=1))

    a = res["A_serie_de_mortes_violentas_intencionais"]
    print("=== A. SERIE DE MORTES VIOLENTAS INTENCIONAIS, BRASIL ===")
    print("   alcance: %d a %d, ou %d anos" % (a["alcance"]["de"], a["alcance"]["a"], a["alcance"]["anos"]))
    print("   edicoes do Anuario lidas: %d" % len(edicoes))
    print("   anos com valor RETIFICADO entre edicoes: %d de %d"
          % (len(retificados), len(anos)))
    print()
    print("   %-6s %-34s %s" % ("ano", "valores publicados", "edicoes"))
    for ano in anos:
        v = a["por_ano"][str(ano)]
        print("   %-6d %-34s %d" % (ano, ", ".join(str(x) for x in v["valores_publicados"]),
                                    v["edicoes_que_publicam"]))
    print()
    for n in notas:
        print("   nota da fonte: " + n)
    print()
    b = res["B_percepcao_de_inseguranca"]
    print("=== B. PERCEPCAO DE INSEGURANCA, 2017 contra 2022 ===")
    print("   estudos lidos: %d" % len(lidos))
    print("   %-52s %8s %8s %s" % ("medo de", "2017", "2022", "direcao"))
    for x in b["comparacao_2017_contra_2022"]:
        d = "sobe" if x["em_2022"] > x["em_2017"] else ("desce" if x["em_2022"] < x["em_2017"] else "mantem")
        print("   %-52s %7.1f%% %7.1f%% %s" % (x["medo_de"], x["em_2017"], x["em_2022"], d))
    print()
    print("   medidas de 2022 sem par em 2017:")
    for x in b["medidas_de_2022_sem_par_em_2017"]:
        print("      %-56s %5.1f%%" % (x["o_que"], x["em_2022"]))
    print()
    print("   a fonte afirma que a queda da violencia NAO foi percebida: %s"
          % ("sim" if b["a_fonte_afirma_que_a_queda_da_violencia_nao_foi_percebida"] else "NAO ACHADO"))
    print()
    print("gravado: " + SAIDA)
    if nao_lidos:
        print()
        print("ATENCAO: %d medida(s) declarada(s) e NAO lida(s) do texto." % len(nao_lidos))
        for x in nao_lidos:
            print("   " + x)
        print("Isso e' ausencia de medicao, e nao ausencia do dado.")
        return 1
    return 0


def bancada():
    provas = []
    t = ("O medo de morrer assassinado cresceu consideravelmente: 82,5% tem este medo em "
         "2022; em 2017 eram 74,9%. Embora tenha sido observada uma queda nas mortes "
         "violentas intencionais, ela nao foi percebida pelos entrevistados.")
    m = re.search(PARES[0]["padrao"], t, re.I | re.S)
    provas.append(("le o par 2017 contra 2022", m is not None))
    provas.append(("o valor de 2022 e' 82,5", m and m.group(1) == "82,5"))
    provas.append(("o valor de 2017 e' 74,9", m and m.group(2) == "74,9"))
    provas.append(("acha a frase da queda nao percebida",
                   re.search(FRASE_CHAVE, re.sub(r"\s+", " ", t), re.I) is not None))
    provas.append(("texto sem o par nao devolve medida",
                   re.search(PARES[0]["padrao"], "texto qualquer", re.I) is None))
    # a direcao tem que ser lida do numero, e nao suposta
    provas.append(("82,5 contra 74,9 e' subida", 82.5 > 74.9))
    provas.append(("o acervo existe no disco", os.path.isdir(ACERVO)))
    ok = sum(1 for _, v in provas if v)
    for nome, v in provas:
        print(("   ok    " if v else "  FALHA  ") + nome)
    print("")
    print("BANCADA DA PERCEPCAO E DA SERIE: %d de %d" % (ok, len(provas)))
    return 0 if ok == len(provas) else 1


if __name__ == "__main__":
    sys.exit(bancada() if "--bancada" in sys.argv else apurar())
